"""C-03 — editing a Word document in place, keeping its look.

``write_document`` builds a NEW .docx from markdown, so "change the fee in this
engagement letter" came back as a retyped document with the letterhead gone.
These tests drive :mod:`iron_jarvis.documents.docx_edit` on a real .docx and
assert the two things the feature is FOR:

* THE FORMATTING SURVIVES — a bold figure is still bold, the header is still
  there, the style of a replaced run is unchanged. Asserting only the text
  would pass just as well for the retyping this module exists to replace.
* AN UNMATCHED EDIT IS A FAILURE, and nothing is written. This is the honesty
  half: a model told "done" over an unchanged file relays that to the user.

The fixture is a fictional engagement letter — never a real client document.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from iron_jarvis.documents.docx_edit import (
    MAX_OPERATIONS,
    DocxEditError,
    edit_docx,
)


# ------------------------------------------------------------------ fixture --


def _letter(path: Path) -> Path:
    """A fictional engagement letter with the things that must survive an edit:
    a heading, a bold figure mid-sentence, a table, and a header."""
    import docx

    doc = docx.Document()
    doc.add_heading("Engagement Letter", level=1)

    par = doc.add_paragraph("Our fee for the 2025 return is ")
    run = par.add_run("$1,250")
    run.bold = True
    par.add_run(" due on filing.")

    doc.add_paragraph("Prepared by Blue Harbor Advisors.")

    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Fee"
    table.cell(1, 0).text = "Form 1040"
    table.cell(1, 1).text = "$900"
    table.cell(2, 0).text = "Form 1120S"
    table.cell(2, 1).text = "$350"

    doc.sections[0].header.paragraphs[0].text = "Northwind Consulting LLC"
    doc.save(str(path))
    return path


def _text(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def _runs_with(path: Path, needle: str) -> list:
    import docx

    doc = docx.Document(str(path))
    return [r for p in doc.paragraphs for r in p.runs if needle in (r.text or "")]


# --------------------------------------------------- 1. the look survives ----


def test_a_replacement_keeps_the_bold_and_the_letterhead(tmp_path):
    """THE WHOLE POINT OF THE ITEM. The figure was bold and mid-paragraph; after
    the edit it is still bold, the heading and header are untouched, and only the
    number changed."""
    src = _letter(tmp_path / "engagement.docx")
    dst = tmp_path / "engagement (edited).docx"

    lines = edit_docx(src, dst, [{"op": "replace", "find": "$1,250", "replace": "$3,000"}])

    assert lines == ['replaced "$1,250" with "$3,000" (1×)']
    body = _text(dst)
    assert "$3,000" in body and "$1,250" not in body
    # The things a retyped document loses:
    assert "Engagement Letter" in body
    assert "Northwind Consulting LLC" in body, "the header did not survive the edit"
    assert "due on filing." in body, "the rest of the sentence was rewritten"

    bold = _runs_with(dst, "$3,000")
    assert bold, "the replacement did not land in a run of its own paragraph"
    assert bold[0].bold is True, (
        "the figure was bold before the edit and is not bold after — this is the "
        "retyping the module exists to avoid, one run smaller"
    )
    # ...and the source is untouched: the default is a copy.
    assert "$1,250" in _text(src)


def test_the_source_is_left_alone_unless_the_destination_IS_the_source(tmp_path):
    """`in_place` is the caller's decision, and the engine honours it literally:
    the same path in and out edits the file, a different one leaves the original
    exactly as it was."""
    src = _letter(tmp_path / "letter.docx")
    before = src.read_bytes()

    edit_docx(src, tmp_path / "copy.docx", [
        {"op": "replace", "find": "Form 1120S", "replace": "Form 1065"}
    ])
    assert src.read_bytes() == before, "a copy-edit rewrote the source"

    edit_docx(src, src, [{"op": "replace", "find": "Form 1120S", "replace": "Form 1065"}])
    assert "Form 1065" in _text(src) and "Form 1120S" not in _text(src)


# ------------------------------------------------------- 2. the honesty ------


def test_an_unmatched_replacement_FAILS_and_names_what_was_missing(tmp_path):
    """Never "Done" over an unchanged file. The error names the text that was not
    found and the file it was looked for in, because that is what a user needs to
    correct the request."""
    src = _letter(tmp_path / "engagement.docx")
    dst = tmp_path / "out.docx"

    with pytest.raises(DocxEditError) as err:
        edit_docx(src, dst, [{"op": "replace", "find": "$9,999", "replace": "$1"}])

    msg = str(err.value)
    assert "$9,999" in msg and "engagement.docx" in msg
    assert "not found" in msg
    assert not dst.exists(), "a failed edit still wrote a file"


def test_a_batch_with_one_bad_operation_writes_NOTHING(tmp_path):
    """All-or-nothing, and it is asserted on the operation ORDER that can only
    pass if the document is applied in memory first: a VALID replace followed by
    a doomed one. A save-as-you-go implementation would leave the first edit on
    disk and call the batch failed."""
    src = _letter(tmp_path / "engagement.docx")
    dst = tmp_path / "out.docx"
    before = src.read_bytes()

    with pytest.raises(DocxEditError):
        edit_docx(src, dst, [
            {"op": "replace", "find": "$1,250", "replace": "$3,000"},   # would work
            {"op": "replace", "find": "nowhere in this file", "replace": "x"},
        ])

    assert not dst.exists()
    assert src.read_bytes() == before, "the source was modified by a failed batch"

    # Same, in place: the file the user already has must not be half-edited.
    with pytest.raises(DocxEditError):
        edit_docx(src, src, [
            {"op": "replace", "find": "$1,250", "replace": "$3,000"},
            {"op": "set_cell", "table": 9, "row": 1, "column": 1, "value": "x"},
        ])
    assert src.read_bytes() == before


def test_the_guards_answer_in_words_a_model_can_act_on(tmp_path):
    src = _letter(tmp_path / "e.docx")
    dst = tmp_path / "o.docx"

    with pytest.raises(DocxEditError, match="at least one operation"):
        edit_docx(src, dst, [])
    with pytest.raises(DocxEditError, match="unknown op"):
        edit_docx(src, dst, [{"op": "rewrite_everything"}])
    with pytest.raises(DocxEditError, match=str(MAX_OPERATIONS)):
        edit_docx(src, dst, [{"op": "replace", "find": "a", "replace": "b"}] * (MAX_OPERATIONS + 1))
    with pytest.raises(DocxEditError, match="`find`"):
        edit_docx(src, dst, [{"op": "replace", "replace": "b"}])
    assert not dst.exists()


# ------------------------------------------------------- 3. the other ops ----


def test_set_cell_is_ONE_BASED_and_says_so_when_it_is_out_of_range(tmp_path):
    """1-based because the request is "row 2, column 2" in the user's words, not
    an index. An out-of-range ask reports what the document actually has."""
    src = _letter(tmp_path / "e.docx")
    dst = tmp_path / "o.docx"

    lines = edit_docx(src, dst, [
        {"op": "set_cell", "table": 1, "row": 2, "column": 2, "value": "$1,100"}
    ])
    assert lines == ['set table 1, row 2, column 2 to "$1,100"']

    import docx

    cell = docx.Document(str(dst)).tables[0].cell(1, 1)
    assert cell.text == "$1,100"
    # The header row and the other row are untouched.
    body = _text(dst)
    assert "Service" in body and "$350" in body

    with pytest.raises(DocxEditError, match="has 3 row\\(s\\), not 9"):
        edit_docx(src, dst, [{"op": "set_cell", "table": 1, "row": 9, "column": 1, "value": "x"}])
    with pytest.raises(DocxEditError, match="the document has 1 table"):
        edit_docx(src, dst, [{"op": "set_cell", "table": 4, "row": 1, "column": 1, "value": "x"}])


def test_insert_after_does_not_turn_the_new_text_into_a_HEADING(tmp_path):
    """Inserting after a heading must produce body text. Inheriting the anchor's
    own style would make every "add a line after the title" request a second
    title — visible to the user, and exactly the kind of look-breaking this item
    is about."""
    src = _letter(tmp_path / "e.docx")
    dst = tmp_path / "o.docx"

    lines = edit_docx(src, dst, [{
        "op": "insert_after",
        "anchor": "Engagement Letter",
        "text": "Revised 11 September 2026.",
    }])
    assert "inserted 1 paragraph(s)" in lines[0]

    import docx

    doc = docx.Document(str(dst))
    idx = next(i for i, p in enumerate(doc.paragraphs) if "Revised 11 September" in p.text)
    style = doc.paragraphs[idx].style.name
    assert not style.lower().startswith("heading") and style.lower() != "title", (
        f"the inserted line was styled {style!r} — it would render as a heading"
    )
    # It really did land AFTER the anchor, not at the end.
    assert "Engagement Letter" in doc.paragraphs[idx - 1].text

    with pytest.raises(DocxEditError, match="no paragraph contains"):
        edit_docx(src, dst, [{"op": "insert_after", "anchor": "No Such Line", "text": "x"}])


def test_a_multi_line_insert_becomes_real_paragraphs(tmp_path):
    src = _letter(tmp_path / "e.docx")
    dst = tmp_path / "o.docx"
    lines = edit_docx(src, dst, [{
        "op": "insert_after",
        "anchor": "Prepared by",
        "text": "First added line.\nSecond added line.",
    }])
    assert "inserted 2 paragraph(s)" in lines[0]

    import docx

    texts = [p.text for p in docx.Document(str(dst)).paragraphs]
    i = texts.index("First added line.")
    assert texts[i + 1] == "Second added line.", "the two lines collapsed into one paragraph"


def test_set_header_footer_reaches_a_part_the_document_has_not_defined(tmp_path):
    """The footer of this fixture is linked to the previous section (it has no
    definition of its own), which is the normal state of a document nobody has
    put a footer in. Setting it must give it one rather than silently writing
    into a part that is not shown."""
    src = _letter(tmp_path / "e.docx")
    dst = tmp_path / "o.docx"

    lines = edit_docx(src, dst, [
        {"op": "set_header_footer", "part": "footer", "text": "Page 1 of 1 — Northwind"},
        {"op": "set_header_footer", "part": "header", "text": "Blue Harbor Advisors"},
    ])
    assert any("footer" in ln for ln in lines) and any("header" in ln for ln in lines)

    import docx

    doc = docx.Document(str(dst))
    footer = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    header = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "Page 1 of 1" in footer
    assert header.strip() == "Blue Harbor Advisors"
    assert doc.sections[0].footer.is_linked_to_previous is False

    with pytest.raises(DocxEditError, match="`part` must be header or footer"):
        edit_docx(src, dst, [{"op": "set_header_footer", "part": "sidebar", "text": "x"}])


# ------------------------------------------- 4. the traversal's two traps ----


def test_a_replacement_reaches_table_cells_and_the_header(tmp_path):
    src = _letter(tmp_path / "e.docx")
    dst = tmp_path / "o.docx"

    edit_docx(src, dst, [{"op": "replace", "find": "Form 1040", "replace": "Form 1040-SR"}])
    assert "Form 1040-SR" in _text(dst)

    edit_docx(src, dst, [{"op": "replace", "find": "Northwind", "replace": "Cedar Point"}])
    import docx

    header = "\n".join(p.text for p in docx.Document(str(dst)).sections[0].header.paragraphs)
    assert "Cedar Point" in header, "the header was never searched"


def test_a_MERGED_cell_is_replaced_ONCE(tmp_path):
    """A merged cell appears in ``row.cells`` once per column it spans, so a
    traversal without de-duplication applies the edit twice: "Fee" -> "Fee total"
    becomes "Fee total total". Pinned because the doubling is invisible in any
    test that only asks whether the new text is present."""
    import docx

    src = tmp_path / "merged.docx"
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 2))
    merged.text = "Fee"
    table.cell(1, 0).text = "900"
    doc.save(str(src))

    dst = tmp_path / "o.docx"
    lines = edit_docx(src, dst, [{"op": "replace", "find": "Fee", "replace": "Fee total"}])

    assert lines == ['replaced "Fee" with "Fee total" (1×)'], (
        "the merged cell was visited more than once"
    )
    cell_text = docx.Document(str(dst)).tables[0].cell(0, 0).text
    assert cell_text == "Fee total", f"the replacement was applied twice: {cell_text!r}"


def test_case_and_count_are_honoured(tmp_path):
    import docx

    src = tmp_path / "c.docx"
    doc = docx.Document()
    doc.add_paragraph("fee FEE Fee")
    doc.save(str(src))
    dst = tmp_path / "o.docx"

    # Case-sensitive by default: only the exact spelling is touched.
    lines = edit_docx(src, dst, [{"op": "replace", "find": "FEE", "replace": "COST"}])
    assert lines == ['replaced "FEE" with "COST" (1×)']
    assert docx.Document(str(dst)).paragraphs[0].text == "fee COST Fee"

    # ...and case-insensitively, all three, unless a count says otherwise.
    lines = edit_docx(src, dst, [
        {"op": "replace", "find": "fee", "replace": "cost", "match_case": False}
    ])
    assert lines == ['replaced "fee" with "cost" (3×)']

    lines = edit_docx(src, dst, [
        {"op": "replace", "find": "fee", "replace": "cost", "match_case": False, "count": 2}
    ])
    assert lines == ['replaced "fee" with "cost" (2×)']
    assert docx.Document(str(dst)).paragraphs[0].text == "cost cost Fee"
