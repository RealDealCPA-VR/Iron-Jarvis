"""v1.291 pins for io-01: reading a workbook whose rows have ordinary BLANK cells.

``_read_xlsx`` treated ``cell.value is None`` as "maybe an uncached formula"
and, for every such row, re-scanned the FORMULA workbook from row 1 with
``iter_rows(min_row=r, max_row=r)`` (read_only sheets re-parse the XML on each
call). A blank cell is ``None`` too, so a normal ledger export with one empty
column paid an O(rows^2) re-parse: 1,000 rows took ~45 s instead of ~0.1 s.
Worse, the lookup used ``cells[0].row`` and an unwritten blank in column A is
an ``EmptyCell`` with no ``.row`` -> AttributeError on any blank separator row.

The fix walks the formula worksheet in LOCKSTEP with the value worksheet
(opened once per sheet, lazily, positional indexing). These pins drive the
real ``extract_text`` -> ``_read_xlsx`` path with real openpyxl workbooks.

Timing is a RATIO on the same machine (blank column vs no blanks, same rows),
never an absolute wall-clock threshold (CLAUDE.md hard rule).
"""

from __future__ import annotations

import time

from openpyxl import Workbook

from iron_jarvis.documents.readers import extract_text

ROWS = 300
COLS = 8


def _make(path, *, blank: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append([f"H{c}" for c in range(COLS)])
    for r in range(ROWS):
        row: list = [f"v{r}-{c}" for c in range(COLS)]
        if blank:
            row[3] = None  # e.g. an optional "Memo" column left empty
        ws.append(row)
    wb.save(path)


def _best_of(path, n: int = 3) -> tuple[float, str]:
    """Fastest of ``n`` reads: damps OneDrive/agent load bursts on this PC."""
    best = float("inf")
    text = ""
    for _ in range(n):
        t0 = time.perf_counter()
        text = extract_text(path)
        best = min(best, time.perf_counter() - t0)
    return best, text


def test_blank_column_is_not_quadratic(tmp_path):
    full = tmp_path / "full.xlsx"
    gaps = tmp_path / "gaps.xlsx"
    _make(full, blank=False)
    _make(gaps, blank=True)
    t_full, text_full = _best_of(full)
    t_gaps, text_gaps = _best_of(gaps)
    ratio = t_gaps / t_full
    # Same row count either way: the blank column must not drop or add rows.
    assert text_gaps.count("\n") == text_full.count("\n")
    assert text_gaps.startswith("## Ledger\n")
    # The blank sheet holds LESS data. The fix costs at most one extra parse
    # (the lockstep formula pass), so ~2x; the defect measured 100x+ at 300
    # rows (592x at 1,000). A ratio, so it holds on any machine.
    assert ratio < 3.0, (
        f"a blank column made the read {ratio:.1f}x slower "
        f"(no blanks {t_full:.3f}s, one blank col {t_gaps:.3f}s)"
    )


def test_blank_row_and_empty_column_a_read_without_crashing(tmp_path):
    """The EmptyCell crash: a blank separator row and a row with column A empty.

    openpyxl (like Excel for unstyled blanks) writes NEITHER to the XML, so the
    read_only reader pads them with EmptyCell, which has no ``.row``.
    """
    path = tmp_path / "gaps.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Fees"
    ws.append(["Client", "Fee", "Memo"])
    ws.append(["Acme", 100, "retainer"])
    ws.append([])  # blank separator row (row 3 has no XML at all)
    ws.append([None, 250, "title in column B only"])  # column A unwritten
    ws.append(["Zed", 75, None])  # trailing blank cell
    wb.save(path)

    text = extract_text(path)

    lines = text.split("\n")
    assert lines[0] == "## Fees"
    assert lines[1] == "Client\tFee\tMemo"
    assert lines[2] == "Acme\t100\tretainer"
    assert lines[3] == "\t\t"  # the blank row is kept, as empty cells
    assert lines[4] == "\t250\ttitle in column B only"
    assert lines[5] == "Zed\t75\t"
    assert len(lines) == 6  # header line + 5 rows, nothing dropped or doubled


def test_uncached_formula_still_shows_formula_text_after_blanks(tmp_path):
    """The rule 'show the formula text where the cached value is missing' must
    survive the lockstep rewrite, including on rows AFTER blank rows (the
    formula pass has to stay aligned with the value pass)."""
    path = tmp_path / "calc.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws.append(["a", "b", "sum"])
    ws.append([1, 2, "=A2+B2"])
    ws.append([])  # blank row shifts the formula pass if it is not in lockstep
    ws.append([None, 4, "=SUM(A4:B4)"])  # column A empty AND a formula
    ws.append([5, 6, None])
    ws.append([7, 8, "=A6*B6"])
    # Second sheet with no blanks at all: must never open the formula pass,
    # and must render identically.
    ws2 = wb.create_sheet("Plain")
    ws2.append(["x", "y"])
    ws2.append([1, 2])
    wb.save(path)

    text = extract_text(path)

    lines = text.split("\n")
    assert lines[0] == "## Calc"
    assert lines[2] == "1\t2\t=A2+B2"
    assert lines[3] == "\t\t"
    assert lines[4] == "\t4\t=SUM(A4:B4)"
    assert lines[5] == "5\t6\t"
    assert lines[6] == "7\t8\t=A6*B6"
    assert lines[7] == "## Plain"
    assert lines[8] == "x\ty"
    assert lines[9] == "1\t2"
    assert len(lines) == 10
