"""v1.244.0 — a spreadsheet a model sends as TEXT is written as the table it
spells, never as one cell.

Replayed on the live model with the v1.244.0 conversation folder in place: the
chat did the job itself in one turn — and the workbook it saved was ONE CELL.
It called write_document with the rows as a JSON string
('[["Date", "Vendor", "Category", "Amount"], ...]'), the .xlsx writer treats a
string as lines, and the whole table landed in A1. The reply said "Done!" over
it and the QA lint called it clean. The tool's schema invites exactly that
slip (content may be "a string, a list of rows, or {'sheets': ...}"), and
local models stringify structured arguments.

Pinned here: JSON rows, JSON sheets, a list of records and a markdown pipe
table all become real rows in .xlsx and .csv; plain text keeps its old one
line per row; nothing a model wrote is dropped.
"""

from __future__ import annotations

import csv
import json

import pytest
from openpyxl import load_workbook

from iron_jarvis.documents.writers import write_document

# The exact shape the live model sent (abridged).
LIVE_JSON_ROWS = json.dumps(
    [
        ["Date", "Vendor", "Category", "Amount"],
        ["2025-01-06", "Henry Schein", "Dental supplies", 1842.17],
        ["2025-01-20", "Patterson Dental", "Dental supplies", 955.00],
        ["", "", "Subtotal — Dental supplies", 2797.17],
        ["", "", "GRAND TOTAL", 2797.17],
    ]
)


def _rows(path) -> list[list]:
    ws = load_workbook(path).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_json_rows_sent_as_a_string_become_real_rows(tmp_path):
    p = write_document(tmp_path / "expenses.xlsx", LIVE_JSON_ROWS)
    rows = _rows(p)
    assert len(rows) == 5, rows  # was 1 — the whole table in A1
    assert rows[0] == ["Date", "Vendor", "Category", "Amount"]
    assert rows[1][1] == "Henry Schein"
    assert rows[1][3] == pytest.approx(1842.17)  # a NUMBER, not text
    assert rows[-1][2:] == ["GRAND TOTAL", pytest.approx(2797.17)]


def test_json_sheets_sent_as_a_string_become_sheets(tmp_path):
    content = json.dumps(
        {"sheets": {"By category": [["Category", "Total"], ["Rent", 12600]], "Raw": [["a", "b"], [1, 2]]}}
    )
    wb = load_workbook(write_document(tmp_path / "book.xlsx", content))
    assert wb.sheetnames == ["By category", "Raw"]
    assert [list(r) for r in wb["By category"].iter_rows(values_only=True)] == [
        ["Category", "Total"],
        ["Rent", 12600],
    ]


def test_sheets_without_the_wrapper_are_still_sheets(tmp_path):
    content = {"Summary": [["Category", "Total"], ["Rent", 12600]], "Detail": [["x"], ["y"]]}
    wb = load_workbook(write_document(tmp_path / "book.xlsx", content))
    assert wb.sheetnames == ["Summary", "Detail"]


def test_records_become_a_header_and_rows(tmp_path):
    records = [
        {"Date": "2025-01-06", "Vendor": "Henry Schein", "Amount": 1842.17},
        {"Date": "2025-01-14", "Vendor": "FPL", "Amount": 612.4, "Note": "late"},
    ]
    for content in (records, json.dumps(records)):
        rows = _rows(write_document(tmp_path / "r.xlsx", content))
        assert rows[0] == ["Date", "Vendor", "Amount", "Note"]
        assert rows[1][1] == "Henry Schein"
        assert rows[2][3] == "late"


def test_a_markdown_table_becomes_rows_and_nothing_around_it_is_lost(tmp_path):
    content = (
        "# Q1 expenses\n\n"
        "| Category | Subtotal |\n|---|---|\n| **Rent** | 12600 |\n| Utilities | 1391.32 |\n\n"
        "Totals include every receipt."
    )
    rows = _rows(write_document(tmp_path / "md.xlsx", content))
    flat = [c for r in rows for c in r if c not in (None, "")]
    assert ["Category", "Subtotal"] in rows
    assert ["Rent", 12600] in rows  # formatting stripped, number kept a number
    assert "Q1 expenses" in flat
    assert "Totals include every receipt." in flat


def test_plain_text_keeps_one_line_per_row(tmp_path):
    rows = _rows(write_document(tmp_path / "t.xlsx", "first line\nsecond line"))
    assert rows == [["first line"], ["second line"]]


def test_a_string_that_merely_contains_brackets_is_left_alone(tmp_path):
    rows = _rows(write_document(tmp_path / "t.xlsx", "Refs [1, 2] and [3]"))
    assert rows == [["Refs [1, 2] and [3]"]]


def test_csv_gets_the_same_recovery(tmp_path):
    p = write_document(tmp_path / "expenses.csv", LIVE_JSON_ROWS)
    with open(p, newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == ["Date", "Vendor", "Category", "Amount"]
    assert len(rows) == 5


def test_real_rows_are_untouched(tmp_path):
    rows = [["a", "b"], [1, 2]]
    assert _rows(write_document(tmp_path / "x.xlsx", rows)) == [["a", "b"], [1, 2]]
