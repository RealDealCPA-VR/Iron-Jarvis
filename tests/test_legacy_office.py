"""C-10 — opening old .xls and .doc files directly.

The fixtures are REAL legacy Office files (committed under ``tests/fixtures``,
authored once by Excel and Word with fictional content), because the whole item
is about the formats themselves: bytes crafted by hand would prove nothing about
a 1997 workbook, and generating them at test time would need Office on the
runner.

Split by what each machine can honestly assert:

* ``.xls`` is read by xlrd — pure Python, so EVERY assertion here runs on CI.
* ``.doc`` needs Word itself, so the conversion tests are Word-gated. The parts
  that must hold everywhere — the honest refusal when Word is absent, and the
  reason a failure reports — are pinned with pure functions and a patched
  availability check, so CI covers them too.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from iron_jarvis.documents import legacy
from iron_jarvis.documents.legacy import (
    LegacyReadError,
    _word_reason,
    read_xls_workbook,
    word_available,
    xls_available,
    xls_text,
)
from iron_jarvis.documents.readers import extract_text

FIXTURES = Path(__file__).parent / "fixtures"
XLS = FIXTURES / "legacy_fees.xls"
DOC = FIXTURES / "legacy_letter.doc"

needs_word = pytest.mark.skipif(
    not word_available(), reason="Microsoft Word is not installed on this machine"
)


def test_the_fixtures_are_really_legacy_office_files():
    """GUARD THE FIXTURES. Both are OLE2 compound files (the 1997 container); if
    one were silently replaced by a modern .docx/.xlsx zip, every test below
    would still pass while testing nothing about legacy formats."""
    for path in (XLS, DOC):
        assert path.is_file(), f"missing fixture: {path}"
        assert path.read_bytes()[:4] == b"\xd0\xcf\x11\xe0", (
            f"{path.name} is not an OLE2 compound file any more"
        )


# ---------------------------------------------------------------- 1. .xls ----


def test_xlrd_is_available_at_all():
    assert xls_available() is True, (
        "xlrd is missing, so .xls support is dead in this build — it is a "
        "pure-Python dependency and belongs in pyproject.toml"
    )


def test_an_old_workbook_is_read_as_TEXT_in_the_same_shape_as_a_modern_one(tmp_path):
    text = xls_text(XLS)

    assert "## Fees" in text, "the sheet heading convention differs from _read_xlsx"
    assert "Northwind Consulting LLC" in text
    assert "Cedar Point Dental" in text
    # Tab-separated rows, like the .xlsx reader.
    row = next(ln for ln in text.splitlines() if "Northwind" in ln)
    assert row.split("\t")[:2] == ["Northwind Consulting LLC", "1250"]


def test_extract_text_reads_it_instead_of_refusing():
    """The old behaviour was a hard stop: "convert it to .xlsx first". That
    sentence must not come back for a file we can now read."""
    text = extract_text(XLS)
    assert "Northwind Consulting LLC" in text
    assert "convert it" not in text


def test_the_values_come_back_TYPED_not_as_strings():
    """A fee read as the string "1250" cannot be summed, and the request is
    almost always arithmetic."""
    book = read_xls_workbook(XLS, "Fees", None)

    assert book["legacy_xls"] is True
    assert book["sheets"] == ["Fees"]
    rows = book["rows"]
    assert rows[0] == ["Client", "Fee", "Paid"]
    fees = [r[1] for r in rows[1:]]
    assert fees == [1250, 980.5, 2100], fees
    assert all(isinstance(f, (int, float)) for f in fees)
    assert sum(fees) == pytest.approx(4330.5)


def test_a_sheet_that_is_not_there_is_named_with_the_ones_that_are():
    with pytest.raises(LegacyReadError) as err:
        xls_text(XLS, sheet="Quarter 4")
    msg = str(err.value)
    assert "Quarter 4" in msg and "Fees" in msg

    with pytest.raises(LegacyReadError, match="there is no sheet 7"):
        xls_text(XLS, sheet=7)


def test_a_cell_range_is_honoured():
    book = read_xls_workbook(XLS, "Fees", "A2:B3")
    assert book["rows"] == [["Northwind Consulting LLC", 1250],
                            ["Blue Harbor Bakery", 980.5]]


# ---------------------------------------------------- 2. .doc, Word-gated ----


@needs_word
def test_word_converts_a_real_doc_and_the_text_comes_back(tmp_path, monkeypatch):
    """THE REGRESSION PIN FOR THE CLEANUP BUG. `$word.Quit(0)` threw
    NonRefArgumentToRefParameterMsg AFTER the .docx had been written, PowerShell
    exited 1, and the caller reported a perfectly good Word document as
    unreadable (measured: rc 1, output file present). A conversion that produced
    a file must be reported as a success."""
    monkeypatch.setattr(legacy, "_cache_dir", lambda: tmp_path / "cache")
    out = tmp_path / "converted.docx"

    legacy.convert_doc_to_docx(DOC, out)

    assert out.is_file() and out.stat().st_size > 0
    assert out.read_bytes()[:2] == b"PK", "the output is not a real .docx zip"

    text = legacy.doc_text(DOC)
    assert "Engagement Letter" in text
    assert "Northwind Consulting LLC" in text
    assert extract_text(DOC).strip() != ""


@needs_word
def test_the_conversion_is_CACHED_by_content(tmp_path, monkeypatch):
    """Word takes seconds to start. The second read of the same bytes must not
    pay for it again — and must not be a stale answer for DIFFERENT bytes."""
    monkeypatch.setattr(legacy, "_cache_dir", lambda: tmp_path / "cache")
    calls = {"n": 0}
    real = legacy.convert_doc_to_docx

    def counted(src, dst, **kw):
        calls["n"] += 1
        return real(src, dst, **kw)

    monkeypatch.setattr(legacy, "convert_doc_to_docx", counted)

    first = legacy.doc_as_docx(DOC)
    second = legacy.doc_as_docx(DOC)

    assert first == second
    assert calls["n"] == 1, f"Word was started {calls['n']} times for one document"

    # Different bytes, different cache entry — never the first answer again.
    other = tmp_path / "other.doc"
    other.write_bytes(DOC.read_bytes() + b"\x00")
    try:
        third = legacy.doc_as_docx(other)
    except LegacyReadError:
        pass  # Word may reject the altered bytes; the key is that it TRIED
    else:
        assert third != first
    assert calls["n"] == 2


# -------------------------------------- 3. convert: legacy -> modern --------
#
# THESE EXIST BECAUSE THE FEATURE SHIPPED DEAD. `convert_document` gates its
# source on `readers.SUPPORTED_READ`, and C-10 taught the READER to open .xls
# and .doc without adding them to that set — so the branch inside the tool that
# converts them could never run, and "turn this .doc into a .docx" came back
# "cannot read '.doc' — supported source formats: ...". Nothing went red: the
# engines were covered, the tool's own suite never converted a legacy file.
# Found by driving a live daemon (v1.218.0's lesson, again).


def _convert_ctx(platform, tmp_path):
    from iron_jarvis.tools.base import ToolContext

    ws = tmp_path / "ws"
    ws.mkdir(parents=True, exist_ok=True)
    return ws, ToolContext(
        workspace=ws, session_id="legacy", agent_run_id="legacy-r1",
        config=platform.config, event_bus=platform.event_bus, engine=platform.engine,
    )


async def test_convert_document_turns_an_old_workbook_into_a_modern_one(platform, tmp_path):
    """CI-SAFE (xlrd is pure Python), and the one that catches the gate: with
    `.xls` out of `SUPPORTED_READ` this refuses before it starts."""
    from openpyxl import load_workbook

    from iron_jarvis.documents.tools import ConvertDocumentTool

    ws, ctx = _convert_ctx(platform, tmp_path)

    res = await ConvertDocumentTool().execute(
        {"source": str(XLS), "target": "fees.xlsx"}, ctx
    )

    assert res.ok is True, res.error
    out = ws / "fees.xlsx"
    assert out.is_file()
    assert res.data["abs_path"] == str(out.resolve())

    rows = [
        [c.value for c in row]
        for row in load_workbook(str(out)).worksheets[0].iter_rows()
    ]
    assert rows[0] == ["Client", "Fee", "Paid"]
    # REAL CELLS, not a text dump: the fee is still a number one can sum.
    fee = next(r[1] for r in rows if r[0] == "Northwind Consulting LLC")
    assert fee == 1250 and isinstance(fee, (int, float))


@needs_word
async def test_convert_document_turns_an_old_letter_into_a_docx(platform, tmp_path):
    """Word re-saves its own 1997 format faithfully — the paragraphs come across
    as paragraphs, not as one flattened block of extracted text."""
    import docx

    from iron_jarvis.documents.tools import ConvertDocumentTool

    ws, ctx = _convert_ctx(platform, tmp_path)

    res = await ConvertDocumentTool().execute(
        {"source": str(DOC), "target": "letter.docx"}, ctx
    )

    assert res.ok is True, res.error
    out = ws / "letter.docx"
    assert out.is_file() and out.read_bytes()[:2] == b"PK"
    texts = [p.text for p in docx.Document(str(out)).paragraphs if p.text.strip()]
    assert any("Engagement Letter" in t for t in texts)
    assert len(texts) > 1, f"the document arrived as one flattened block: {texts}"


def test_the_legacy_pair_is_ADVERTISED_as_readable():
    """The set every caller gates on. `extract_text` reading a format while
    `SUPPORTED_READ` denies it is not a cosmetic mismatch — it is what refused
    the conversion above."""
    from iron_jarvis.documents.readers import SUPPORTED_READ

    assert ".xls" in SUPPORTED_READ and ".doc" in SUPPORTED_READ


def test_the_unattended_folder_sweep_still_declines_dot_doc():
    """A deliberate boundary, not an oversight: every unique .doc starts Word
    over COM, so a folder of fifty would start Word fifty times on the machine
    the user is working on. Reading one directly is unaffected."""
    from iron_jarvis.documents.batch import SWEEP_SUFFIXES, SWEEP_SUFFIXES_WITH_OCR

    assert ".xls" in SWEEP_SUFFIXES, "a pure-Python workbook has no reason to be skipped"
    assert ".doc" not in SWEEP_SUFFIXES
    assert ".doc" not in SWEEP_SUFFIXES_WITH_OCR


# ------------------------------------- 4. what must hold with NO Word (CI) ---


def test_with_no_word_the_refusal_names_WORD_and_not_a_chore(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy, "_cache_dir", lambda: tmp_path / "cache")
    monkeypatch.setattr(legacy, "word_available", lambda: False)

    with pytest.raises(LegacyReadError) as err:
        legacy.doc_text(DOC)

    msg = str(err.value)
    assert "Microsoft Word is not installed" in msg
    assert ".docx" in msg, "it did not say what would work instead"
    # The old blanket refusal is gone: it never said WHY, only what to do.
    assert "convert it to .docx first" not in msg


def test_the_failure_reason_is_never_the_CLIXML_BANNER():
    """What the user was actually shown before this was fixed: "Microsoft Word
    could not convert letter.doc: #< CLIXML". The stderr of a piped
    `-EncodedCommand` PowerShell is serialised error records, so its first line
    is always that banner — the reason has to be dug out of the XML (or, better,
    read from the marker the script writes on stdout)."""
    clixml = (
        '#< CLIXML\r\n<Objs Version="1.1.0.1" xmlns="http://schemas.microsoft.com/'
        'powershell/2004/04"><S S="Error">Argument: \'1\' should be a '
        "System.Management.Automation.PSReference. Use [ref]._x000D__x000A_</S>"
        '<S S="Error">At line:13 char:3_x000D__x000A_</S></Objs>'
    ).encode("utf-8")

    reason = _word_reason("", clixml)
    # THE MESSAGE MUST BE THE MESSAGE, not the envelope it travelled in. These
    # assertions are deliberately stricter than "CLIXML not in reason": with the
    # extraction disabled, the fallback returns the whole `<Objs ...>` line,
    # which still contains "PSReference" and no "CLIXML" — so a looser test
    # passed while the user would have been shown raw XML (measured: this test's
    # first cut survived exactly that mutation).
    assert reason.startswith("Argument:"), reason
    assert "PSReference" in reason
    assert "<" not in reason and ">" not in reason, f"raw XML reached the reason: {reason}"
    assert "_x000D_" not in reason and "_x000A_" not in reason, reason
    assert "CLIXML" not in reason and "Objs" not in reason
    assert not reason.startswith("At line:"), "it picked the position record"

    # The script's own marker wins over anything on stderr.
    assert _word_reason("IJ-ERROR: the document is password-protected", clixml) == (
        "the document is password-protected"
    )
    # A plain (non-CLIXML) stderr line is used as-is.
    assert _word_reason("", b"Word is busy\n") == "Word is busy"
    # And with nothing to go on, a sentence rather than an empty string.
    assert _word_reason("", b"") == "Word could not open the file"


def test_a_password_protected_document_is_named_as_such(tmp_path, monkeypatch):
    """Word reports a password through its own error text; the reader turns that
    into the one thing the user needs to know."""
    monkeypatch.setattr(legacy, "_cache_dir", lambda: tmp_path / "cache")
    monkeypatch.setattr(legacy, "word_available", lambda: True)

    class FakeProc:
        returncode = 3

        def communicate(self, timeout=None):
            return (b"IJ-ERROR: This document is protected by a password.\n", b"")

    monkeypatch.setattr(legacy.subprocess, "Popen", lambda *a, **k: FakeProc())

    with pytest.raises(LegacyReadError) as err:
        legacy.convert_doc_to_docx(DOC, tmp_path / "x.docx")
    assert "password-protected" in str(err.value)


def test_a_conversion_that_writes_nothing_is_a_FAILURE(tmp_path, monkeypatch):
    """The other direction of the exit-code lesson: the verdict is the marker
    PLUS the file, so a script that claims OK without producing a document must
    not be believed."""
    monkeypatch.setattr(legacy, "word_available", lambda: True)

    class LyingProc:
        returncode = 0

        def communicate(self, timeout=None):
            return (b"IJ-OK\n", b"")

    monkeypatch.setattr(legacy.subprocess, "Popen", lambda *a, **k: LyingProc())

    with pytest.raises(LegacyReadError):
        legacy.convert_doc_to_docx(DOC, tmp_path / "never-written.docx")


def test_a_timeout_is_reported_as_a_timeout(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy, "word_available", lambda: True)
    killed: list[int] = []

    class HangingProc:
        pid = 4242
        returncode = None

        def __init__(self):
            self.calls = 0

        def communicate(self, timeout=None):
            self.calls += 1
            if self.calls == 1:
                raise legacy.subprocess.TimeoutExpired(cmd="powershell", timeout=timeout or 1)
            return (b"", b"")

    monkeypatch.setattr(legacy.subprocess, "Popen", lambda *a, **k: HangingProc())
    monkeypatch.setattr(legacy, "_kill_tree", lambda pid: killed.append(pid))

    with pytest.raises(LegacyReadError) as err:
        legacy.convert_doc_to_docx(DOC, tmp_path / "x.docx", timeout=1)

    assert "did not finish" in str(err.value)
    assert killed == [4242], "the Word process tree was not killed after the timeout"
