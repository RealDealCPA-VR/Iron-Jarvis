"""C-07 — comparing two documents exactly.

"What changed between last year's return and this year's" was answered by a
model reading both files and describing the differences from memory, which is the
one reliable way to miss a changed figure. These tests drive
:mod:`iron_jarvis.documents.compare` and assert the differences it COMPUTES:
the cell references, the numbers, and the deltas.

Every fixture is fictional.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from iron_jarvis.documents import compare as compare_mod
from iron_jarvis.documents.compare import (
    CompareError,
    compare_documents,
    render_report,
)


# ----------------------------------------------------------------- fixtures --


def _letter(path: Path, fee: str, extra: str | None = None) -> Path:
    import docx

    doc = docx.Document()
    doc.add_heading("Engagement Letter", level=1)
    doc.add_paragraph("Client: Northwind Consulting LLC")
    doc.add_paragraph(f"Our fee for the 2025 return is {fee} due on filing.")
    doc.add_paragraph("Returns included: Form 1040, Form 1120S")
    if extra:
        doc.add_paragraph(extra)
    doc.save(str(path))
    return path


def _book(path: Path, rows: list[list], sheet: str = "Fees") -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.save(str(path))
    return path


_ROWS = [
    ["Client", "Fee", "Paid"],
    ["Northwind Consulting LLC", 1250, "yes"],
    ["Blue Harbor Bakery", 980, "no"],
    ["Cedar Point Dental", 2100, "yes"],
]


# ----------------------------------------------------- 1. text: the figures --


def test_a_changed_fee_is_reported_with_the_DELTA(tmp_path):
    """The headline case. Not just "this paragraph changed" — the figure that
    moved, with the arithmetic done, because that is what the user is asking."""
    a = _letter(tmp_path / "2024.docx", "$1,250")
    b = _letter(tmp_path / "2025.docx", "$3,000")

    result = compare_documents(a, b)

    assert result["mode"] == "text"
    assert result["identical"] is False
    assert result["counts"]["changed"] == 1
    figures = result["figures"]
    assert len(figures) == 1, figures
    fig = figures[0]
    assert fig["before"] == "$1,250" and fig["after"] == "$3,000"
    assert fig["change"] == "+1,750", f"the delta was not computed: {fig}"
    assert "fee" in fig["where"].casefold()

    report = render_report(result)
    assert "$1,250" in report and "$3,000" in report and "+1,750" in report


def test_added_and_removed_paragraphs_are_counted_separately(tmp_path):
    a = _letter(tmp_path / "a.docx", "$1,250")
    b = _letter(tmp_path / "b.docx", "$1,250", extra="Late filing fee: $150 per month.")

    result = compare_documents(a, b)

    assert result["counts"]["added"] == 1
    assert result["counts"]["removed"] == 0
    assert result["counts"]["changed"] == 0
    kinds = [c["kind"] for c in result["changes"]]
    assert kinds == ["added"], kinds
    assert "Late filing fee" in render_report(result)


def test_two_identical_documents_say_so_in_words(tmp_path):
    a = _letter(tmp_path / "a.docx", "$1,250")
    b = _letter(tmp_path / "b.docx", "$1,250")

    result = compare_documents(a, b)

    assert result["identical"] is True
    assert result["changes"] == [] and result["figures"] == []
    report = render_report(result)
    assert "no differences" in report and "a.docx" in report and "b.docx" in report


def test_figures_mode_matches_by_LABEL_not_by_position(tmp_path):
    """`figures` answers "what numbers moved" even when the lines were reordered
    — the labels are the key, so a paragraph that moved down the page does not
    read as two changes."""
    import docx

    a = tmp_path / "a.docx"
    doc = docx.Document()
    doc.add_paragraph("Rental income: 12,000")
    doc.add_paragraph("Depreciation: 3,400")
    doc.save(str(a))

    b = tmp_path / "b.docx"
    doc = docx.Document()
    doc.add_paragraph("Depreciation: 3,400")      # same value, moved up
    doc.add_paragraph("Rental income: 15,500")    # changed
    doc.save(str(b))

    result = compare_documents(a, b, mode="figures")

    assert result["mode"] == "figures"
    rows = result["figures"]
    assert len(rows) == 1, rows
    assert rows[0]["where"] == "Rental income"
    assert rows[0]["before"] == "12,000" and rows[0]["after"] == "15,500"
    assert rows[0]["change"] == "+3,500"


# --------------------------------------------------- 2. cells: spreadsheets --


def test_a_spreadsheet_is_compared_CELL_BY_CELL_with_A1_references(tmp_path):
    a = _book(tmp_path / "q1.xlsx", _ROWS)
    changed = [list(r) for r in _ROWS]
    changed[2][1] = 1180  # Blue Harbor Bakery's fee
    b = _book(tmp_path / "q2.xlsx", changed)

    result = compare_documents(a, b)

    assert result["mode"] == "cells"
    assert result["counts"]["cells"] == 1
    cell = result["changes"][0]
    assert cell["cell"] == "B3", f"wrong reference: {cell}"
    assert cell["sheet"] == "Fees"
    assert cell["before"] == "980" and cell["after"] == "1180"
    assert cell["change"] == "+200"
    assert "Fees!B3" in render_report(result)


def test_an_INSERTED_ROW_does_not_make_every_later_row_look_changed(tmp_path):
    """THE REASON `label_column` EXISTS. Positionally, inserting one row shifts
    every row below it and the whole sheet reads as changed — the report is then
    useless for the question actually asked. Matched on the label column, the
    answer is one added row."""
    a = _book(tmp_path / "a.xlsx", _ROWS)
    inserted = [list(r) for r in _ROWS]
    inserted.insert(2, ["Aspen Ridge Veterinary", 640, "no"])
    b = _book(tmp_path / "b.xlsx", inserted)

    positional = compare_documents(a, b)
    assert positional["counts"]["cells"] >= 4, (
        "the positional comparison did not shift — this test no longer contrasts "
        "the two modes"
    )

    matched = compare_documents(a, b, label_column="Client")
    assert matched["counts"]["cells"] == 0, matched["changes"]
    assert matched["counts"]["rows_added"] == 1
    assert matched["counts"]["rows_removed"] == 0
    added = [c for c in matched["changes"] if c["kind"] == "row_added"]
    assert added[0]["row"] == "Aspen Ridge Veterinary"
    assert "only in the second" in render_report(matched)


def test_label_matching_reports_a_removed_row_and_a_changed_field_together(tmp_path):
    a = _book(tmp_path / "a.xlsx", _ROWS)
    edited = [list(r) for r in _ROWS]
    edited.pop(2)              # Blue Harbor Bakery is gone
    edited[1][2] = "no"        # Northwind stopped paying
    b = _book(tmp_path / "b.xlsx", edited)

    result = compare_documents(a, b, label_column="Client")

    assert result["counts"]["rows_removed"] == 1
    removed = [c for c in result["changes"] if c["kind"] == "row_removed"]
    assert removed[0]["row"] == "Blue Harbor Bakery"
    cells = [c for c in result["changes"] if c["kind"] == "cell"]
    assert len(cells) == 1
    assert cells[0]["row"] == "Northwind Consulting LLC"
    assert cells[0]["column"] == "Paid"
    assert cells[0]["before"] == "yes" and cells[0]["after"] == "no"


def test_a_column_that_MOVED_is_still_compared_to_itself(tmp_path):
    """Columns are matched by header name under `label_column`, so swapping two
    columns is not a sheet full of differences."""
    a = _book(tmp_path / "a.xlsx", _ROWS)
    swapped = [[r[0], r[2], r[1]] for r in _ROWS]  # Client, Paid, Fee
    b = _book(tmp_path / "b.xlsx", swapped)

    result = compare_documents(a, b, label_column="Client")

    assert result["counts"]["cells"] == 0, result["changes"]
    assert result["identical"] is True


def test_a_csv_and_a_workbook_are_paired_even_with_different_sheet_names(tmp_path):
    a = tmp_path / "last_year.csv"
    a.write_text(
        "\n".join(",".join(str(c) for c in row) for row in _ROWS),
        encoding="utf-8",
    )
    changed = [list(r) for r in _ROWS]
    changed[1][1] = 1400
    b = _book(tmp_path / "this_year.xlsx", changed, sheet="2025")

    result = compare_documents(a, b)

    assert result["sheets_only_in_first"] == []
    assert result["sheets_only_in_second"] == []
    assert result["counts"]["cells"] == 1
    assert result["changes"][0]["cell"] == "B2"


def test_a_sheet_present_in_only_one_workbook_is_NAMED(tmp_path):
    from openpyxl import Workbook

    a = tmp_path / "a.xlsx"
    wb = Workbook()
    wb.active.title = "Fees"
    wb.create_sheet("Notes")
    wb.save(str(a))

    b = tmp_path / "b.xlsx"
    wb = Workbook()
    wb.active.title = "Fees"
    wb.save(str(b))

    result = compare_documents(a, b)

    assert result["sheets_only_in_first"] == ["Notes"]
    assert result["identical"] is False
    assert "Notes" in render_report(result)


# ------------------------------------------------------ 3. refusals & caps --


def test_the_refusals_name_what_to_do_instead(tmp_path):
    a = _letter(tmp_path / "a.docx", "$1")
    b = _letter(tmp_path / "b.docx", "$2")
    book = _book(tmp_path / "c.xlsx", _ROWS)

    with pytest.raises(CompareError, match="unknown mode"):
        compare_documents(a, b, mode="telepathy")
    with pytest.raises(CompareError, match="needs two spreadsheets"):
        compare_documents(a, b, mode="cells")
    with pytest.raises(CompareError) as err:
        compare_documents(book, book, label_column="Nonexistent Column")
    # The header row is quoted, so the caller can pick a real column.
    assert "Client" in str(err.value) and "Nonexistent Column" in str(err.value)


def test_reading_only_NEVER_touches_either_file(tmp_path):
    a = _book(tmp_path / "a.xlsx", _ROWS)
    changed = [list(r) for r in _ROWS]
    changed[1][1] = 9
    b = _book(tmp_path / "b.xlsx", changed)
    before_a, before_b = a.read_bytes(), b.read_bytes()

    compare_documents(a, b)
    compare_documents(a, b, label_column="Client")

    assert a.read_bytes() == before_a and b.read_bytes() == before_b


def test_a_cap_that_BITES_is_reported(tmp_path, monkeypatch):
    """A truncated answer that reads as complete is worse than a short one: the
    model relays "3 cells changed" over a sheet where thirty did."""
    monkeypatch.setattr(compare_mod, "MAX_CHANGES", 3)
    a = _book(tmp_path / "a.xlsx", [[i, i] for i in range(40)])
    b = _book(tmp_path / "b.xlsx", [[i, i + 1] for i in range(40)])

    result = compare_documents(a, b)

    assert result["truncated"] is True
    assert len(result["changes"]) <= 3
    assert "more" in render_report(result)


def test_the_result_carries_both_paths(tmp_path):
    a = _letter(tmp_path / "a.docx", "$1")
    b = _letter(tmp_path / "b.docx", "$2")
    result = compare_documents(a, b)
    assert result["first"] == str(a) and result["second"] == str(b)
