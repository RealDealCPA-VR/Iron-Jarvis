"""Edit an existing Word document and keep its look (C-03).

``write_document`` builds a NEW .docx from markdown, so "change the fee to
$3,000 in this engagement letter" used to come back as a plain retyped
document — the letterhead, styles, headers and numbering gone. This module edits
the document python-docx OPENED instead:

* ``replace`` swaps exact text run-by-run through the SAME helper redaction uses
  (:func:`.redact._redact_runs`), so a bold figure stays bold and a replacement
  that spans two runs lands in the run where it starts;
* ``set_cell`` / ``insert_after`` / ``set_header_footer`` change one table cell,
  add paragraphs after a named one, or rewrite a header/footer, keeping the
  existing run formatting where there is one.

HONESTY, and it is the point of the module: a replacement that matches nothing
is an ERROR naming the text — never "Done" over an unchanged file — and a batch
with any failing operation writes NOTHING (every operation runs on the in-memory
document first; the file is saved only when all of them succeeded).

Known limits, stated rather than hidden: text inside hyperlinks, tracked
insertions and text boxes is not searched (python-docx does not expose those as
paragraph runs), and a phrase that crosses a paragraph break cannot match.
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any, Iterator

#: Operations per call — a guard against a runaway model loop, far above any
#: real edit request.
MAX_OPERATIONS = 200

OPS = ("replace", "set_cell", "insert_after", "set_header_footer")


class DocxEditError(ValueError):
    """An operation that could not be applied. The message is shown to the
    model (and relayed to the user), so it names what was not found."""


# --------------------------------------------------------------- traversal ---


def _container_paragraphs(container: Any) -> Iterator[Any]:
    """Every paragraph in *container*, tables (and nested tables) included."""
    for par in container.paragraphs:
        yield par
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _container_paragraphs(cell)


_HEADER_FOOTER_ATTRS = (
    "header",
    "footer",
    "first_page_header",
    "first_page_footer",
    "even_page_header",
    "even_page_footer",
)


def _all_paragraphs(doc: Any) -> Iterator[Any]:
    """Body, tables and every header/footer the document really DEFINES, each
    underlying paragraph exactly once.

    Once, because a merged table cell appears in ``row.cells`` for every column
    it spans — the same paragraph visited twice would apply a replacement twice
    (``"Fee"`` → ``"Fee total"`` would become ``"Fee total total"``).

    THE SET HOLDS THE ELEMENTS, NEVER ``id(par._p)``, and that is a correctness
    requirement rather than a style choice. lxml builds a proxy object on each
    ``._p`` access and frees it as soon as the last reference goes, so a later,
    entirely unrelated paragraph's proxy can be allocated at the SAME address and
    then reads as "already seen" — paragraphs were silently dropped from the
    search, and WHICH ones depended on allocator luck. Measured on one fixture:
    a table row ("Form 1120S") vanished in one test and the header vanished in
    another, in the same run, both reported as "not found in <file>". Keeping the
    element in the set holds a reference, which is exactly what makes its
    identity stable for as long as the traversal lasts."""
    seen: set[Any] = set()

    def fresh(pars: Iterator[Any]) -> Iterator[Any]:
        for par in pars:
            el = par._p
            if el in seen:
                continue
            seen.add(el)
            yield par

    yield from fresh(_container_paragraphs(doc))
    for section in doc.sections:
        for attr in _HEADER_FOOTER_ATTRS:
            part = getattr(section, attr, None)
            # A linked header has no definition of its own — its text is the
            # previous section's, already visited.
            if part is None or part.is_linked_to_previous:
                continue
            yield from fresh(_container_paragraphs(part))


# ---------------------------------------------------------------- the ops ---


def _replace(doc: Any, find: str, repl: str, *, match_case: bool, limit: int | None) -> int:
    from .redact import _redact_runs

    rx = re.compile(re.escape(find), 0 if match_case else re.IGNORECASE)
    remaining: list[int | None] = [limit]
    total = 0
    for par in _all_paragraphs(doc):
        if remaining[0] is not None and remaining[0] <= 0:
            break
        runs = list(par.runs)
        if not runs:
            continue
        if not rx.search("".join(r.text or "" for r in runs)):
            continue

        def spans_for(text: str) -> list[tuple[int, int, str, str]]:
            out: list[tuple[int, int, str, str]] = []
            for m in rx.finditer(text):
                if remaining[0] is not None:
                    if remaining[0] <= 0:
                        break
                    remaining[0] -= 1
                out.append((m.start(), m.end(), "replace", repl))
            return out

        total += _redact_runs(runs, spans_for).get("replace", 0)
    return total


def _set_container_text(container: Any, text: str) -> None:
    """Replace everything in *container* (a cell, header or footer) with *text*,
    keeping the FIRST run's formatting — the font a cell or letterhead already
    uses is the one the new text should wear."""
    pars = list(container.paragraphs)
    if not pars:
        container.add_paragraph(text)
        return
    first = pars[0]
    runs = list(first.runs)
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        first.add_run(text)
    for extra in pars[1:]:
        el = extra._p
        el.getparent().remove(el)


def _set_cell(doc: Any, table: int, row: int, column: int, value: str) -> str:
    tables = doc.tables
    if not 1 <= table <= len(tables):
        raise DocxEditError(
            f"there is no table {table} — the document has {len(tables)} table(s)"
        )
    t = tables[table - 1]
    if not 1 <= row <= len(t.rows):
        raise DocxEditError(f"table {table} has {len(t.rows)} row(s), not {row}")
    cells = t.rows[row - 1].cells
    if not 1 <= column <= len(cells):
        raise DocxEditError(
            f"row {row} of table {table} has {len(cells)} column(s), not {column}"
        )
    _set_container_text(cells[column - 1], value)
    return f"set table {table}, row {row}, column {column} to \"{value}\""


def _is_heading(style_name: str) -> bool:
    s = (style_name or "").strip().lower()
    return s.startswith("heading") or s in ("title", "subtitle")


def _insert_after(doc: Any, anchor: str, text: str, style: str | None) -> str:
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    body = list(doc.paragraphs)
    idx = next((i for i, p in enumerate(body) if anchor in p.text), None)
    if idx is None:  # a second, case-insensitive look before giving up
        low = anchor.casefold()
        idx = next((i for i, p in enumerate(body) if low in p.text.casefold()), None)
    if idx is None:
        raise DocxEditError(
            f"no paragraph contains \"{anchor}\" — nothing was inserted"
        )
    target = body[idx]
    if style:
        style_name = style
    else:
        own = target.style.name if target.style is not None else "Normal"
        if not _is_heading(own):
            style_name = own
        else:
            # After a heading, new text should look like the text that follows
            # it, not like another heading.
            nxt = body[idx + 1] if idx + 1 < len(body) else None
            nxt_style = nxt.style.name if nxt is not None and nxt.style is not None else ""
            style_name = nxt_style if nxt_style and not _is_heading(nxt_style) else "Normal"
    last = target
    lines = text.split("\n")
    for line in lines:
        new_p = OxmlElement("w:p")
        last._p.addnext(new_p)
        para = Paragraph(new_p, last._parent)
        try:
            para.style = style_name
        except KeyError as exc:
            raise DocxEditError(f"this document has no style named \"{style_name}\"") from exc
        para.add_run(line)
        last = para
    return f"inserted {len(lines)} paragraph(s) after \"{anchor}\""


def _set_header_footer(doc: Any, part: str, text: str, section: int | None) -> str:
    sections = list(doc.sections)
    if section is not None and not 1 <= section <= len(sections):
        raise DocxEditError(
            f"there is no section {section} — the document has {len(sections)}"
        )
    indexes = [section - 1] if section is not None else range(len(sections))
    changed = 0
    for i in indexes:
        sec = sections[i]
        hf = sec.header if part == "header" else sec.footer
        if hf.is_linked_to_previous:
            if section is None and i > 0:
                continue  # it shows the previous section's, already set
            hf.is_linked_to_previous = False  # give it a definition of its own
        _set_container_text(hf, text)
        changed += 1
    where = f"section {section}" if section is not None else f"{changed} section(s)"
    return f"set the {part} of {where} to \"{text}\""


# ------------------------------------------------------------------ entry ---


def _as_int(value: Any, what: str, op_no: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise DocxEditError(f"operation {op_no}: `{what}` must be a number") from exc


def edit_docx(src: Path, dst: Path, operations: list[dict[str, Any]]) -> list[str]:
    """Apply *operations* to *src* and save to *dst* (which may be *src*).

    Returns one plain line per operation. Raises :class:`DocxEditError` — with
    NOTHING written — when any operation cannot be applied."""
    import docx

    from .readers import _guard_office_encrypted

    if not operations:
        raise DocxEditError("nothing to do — pass at least one operation")
    if len(operations) > MAX_OPERATIONS:
        raise DocxEditError(f"too many operations ({len(operations)}); the limit is {MAX_OPERATIONS}")
    _guard_office_encrypted(src)
    doc = docx.Document(str(src))
    lines: list[str] = []
    for n, op in enumerate(operations, 1):
        if not isinstance(op, dict):
            raise DocxEditError(f"operation {n} is not an object")
        kind = str(op.get("op") or "").strip().lower()
        if kind == "replace":
            find = str(op.get("find") or "")
            if not find:
                raise DocxEditError(f"operation {n}: replace needs the `find` text")
            repl = "" if op.get("replace") is None else str(op.get("replace"))
            raw_count = op.get("count")
            limit = None if raw_count in (None, "", "all") else _as_int(raw_count, "count", n)
            got = _replace(
                doc, find, repl,
                match_case=op.get("match_case", True) is not False,
                limit=limit,
            )
            if got == 0:
                raise DocxEditError(
                    f"operation {n}: \"{find}\" was not found in {src.name} — "
                    "nothing was changed or saved"
                )
            lines.append(f"replaced \"{find}\" with \"{repl}\" ({got}×)")
        elif kind == "set_cell":
            lines.append(_set_cell(
                doc,
                _as_int(op.get("table", 1), "table", n),
                _as_int(op.get("row"), "row", n),
                _as_int(op.get("column"), "column", n),
                "" if op.get("value") is None else str(op.get("value")),
            ))
        elif kind == "insert_after":
            anchor = str(op.get("anchor") or "")
            if not anchor:
                raise DocxEditError(f"operation {n}: insert_after needs the `anchor` text")
            lines.append(_insert_after(
                doc, anchor, str(op.get("text") or ""),
                str(op.get("style") or "").strip() or None,
            ))
        elif kind == "set_header_footer":
            part = str(op.get("part") or "header").strip().lower()
            if part not in ("header", "footer"):
                raise DocxEditError(f"operation {n}: `part` must be header or footer")
            section = op.get("section")
            lines.append(_set_header_footer(
                doc, part, str(op.get("text") or ""),
                None if section in (None, "", "all") else _as_int(section, "section", n),
            ))
        else:
            raise DocxEditError(
                f"operation {n}: unknown op \"{kind}\" — use one of {', '.join(OPS)}"
            )
    dst.parent.mkdir(parents=True, exist_ok=True)
    tmp = dst.with_name(f".{dst.name}.tmp-{os.getpid()}")
    try:
        doc.save(str(tmp))
        os.replace(tmp, dst)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
    return lines
